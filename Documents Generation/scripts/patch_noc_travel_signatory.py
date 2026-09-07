"""Add the {{signatory_name}} line to noc-travel.docx (tourist-visa maid NOC).

The tourist NOC's signature block was a static two-line company footer:

    Maids CC Domestic Worker Services LLC
    0195567

The global signatory rule requires every generated NOC to carry the configured signatory
(default "HR Manager") plus the signature and company-stamp images. This inserts a
``{{signatory_name}}`` paragraph immediately above the company legal line so the rendered
block reads:

    [signature image]
    HR Manager
    [stamp image]
    Maids CC Domestic Worker Services LLC
    0195567

The company legal name + trade-license number are preserved as the issuing entity. Idempotent:
running it again is a no-op once the token is present.

    python scripts/patch_noc_travel_signatory.py
"""

from pathlib import Path

from docx import Document

BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = BASE_DIR / "noc-travel.docx"

COMPANY_LINE = "Maids CC Domestic Worker Services LLC"
SIGNATORY_TOKEN = "{{signatory_name}}"


def main() -> None:
    if not TEMPLATE.exists():
        raise SystemExit(f"Missing template: {TEMPLATE}")

    doc = Document(TEMPLATE)

    # Already patched?
    for para in doc.paragraphs:
        if SIGNATORY_TOKEN in "".join(run.text for run in para.runs):
            print(f"Already patched -> {TEMPLATE.name}")
            return

    target = None
    for para in doc.paragraphs:
        if "".join(run.text for run in para.runs).strip() == COMPANY_LINE:
            target = para
            break
    if target is None:
        raise SystemExit(f'Could not find signature line "{COMPANY_LINE}" in {TEMPLATE.name}')

    # Insert the signatory paragraph directly above the company legal line, matching its style
    # so the fill flow (which renders substituted values bold) styles the name consistently.
    signatory = target.insert_paragraph_before()
    signatory.alignment = target.alignment
    if target.style is not None:
        signatory.style = target.style
    signatory.add_run(SIGNATORY_TOKEN)

    doc.save(TEMPLATE)
    print(f"Inserted {SIGNATORY_TOKEN} above '{COMPANY_LINE}' -> {TEMPLATE.name}")


if __name__ == "__main__":
    main()
